"""Excel file and the JSON the public page reads."""
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Bid

COLUMNS = [
    ("County", "county", 14), ("Agency", "agency", 30), ("Bid #", "number", 20), ("Title", "title", 60),
    ("Category", "category", 26), ("Type", "type", 18), ("Status", "status", 16), ("Due", "due", 18),
    ("Posted", "posted", 12), ("Bid link", "bid_url", 40), ("Source page", "source_url", 40), ("Notes", "notes", 40),
]


def _due_cell(iso: str):
    if not iso:
        return ""
    if len(iso) == 10:
        return datetime.fromisoformat(iso)
    return datetime.fromisoformat(iso).replace(tzinfo=None)


def write_xlsx(bids: list[Bid], path: Path, updated: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Open solicitations"
    ws.append([c[0] for c in COLUMNS])
    head_fill = PatternFill("solid", fgColor="0E4D45")
    for i, (_, _, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        ws.column_dimensions[get_column_letter(i)].width = width
    for b in bids:
        row = []
        for _, field, _ in COLUMNS:
            v = getattr(b, field)
            if field == "due":
                v = _due_cell(v)
            elif field == "posted" and v:
                v = datetime.fromisoformat(v)
            row.append(v)
        ws.append(row)
        r = ws.max_row
        for col, field in ((10, "bid_url"), (11, "source_url")):
            url = getattr(b, field)
            if url:
                c = ws.cell(row=r, column=col)
                c.hyperlink = url
                c.font = Font(color="0B5CAD", underline="single")
        due = ws.cell(row=r, column=8)
        due.number_format = "yyyy-mm-dd" if len(b.due) == 10 else "yyyy-mm-dd h:mm AM/PM"
        ws.cell(row=r, column=9).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=4).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    info = wb.create_sheet("About")
    info.append(["floridabids.org: open public solicitations in Florida"])
    info.append([f"Updated {updated} (Eastern time)"])
    info.append(["Always confirm details and deadlines on the agency's own page before responding."])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_site_json(bids: list[Bid], meta: dict, path: Path) -> None:
    keep = ["county", "agency", "number", "title", "category", "type", "status", "due", "posted",
            "bid_url", "source_url", "notes", "first_seen"]
    rows = [{k: getattr(b, k) for k in keep} for b in bids]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps({"meta": meta, "bids": rows}, ensure_ascii=False, separators=(",", ":")))
